#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
pjt_fill_buffer.py

把 buffer 工作表的 pd-dz 欄（col D）改成指向指定週次 sheet 的公式。
例：='07-14'!K5  或  ='07-14'!AB35

週次 sheet 有兩個面板：
  左面板：col 1 = section header（IND/CAB/SOLO），月份在 col 5-13
  右面板：col 18 = section header（CHN/NVN），月份在 col 22-30

用法（dry-run 先確認）：
  python pjt_fill_buffer.py <pjt.xlsx> --week 07-14 --dry-run
確認後正式寫入：
  python pjt_fill_buffer.py <pjt.xlsx> --week 07-14

直接修改 xlsx，執行前自動備份。
"""

import sys
import re
import argparse
import shutil
from pathlib import Path
import openpyxl
import openpyxl.reader.excel as _xlreader
from openpyxl.utils import get_column_letter

# 容忍 PJT 大表壞 drawing 參照
_orig = _xlreader.find_images
def _safe(a, p):
    try:
        return _orig(a, p)
    except Exception:
        return [], []
_xlreader.find_images = _safe

REGIONS = ['IND', 'CAB', 'CHN', 'NVN']
MONTH_ABBRS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
               'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

# 左面板（IND/CAB/SOLO）的 section header 在 col 1
LEFT_HEADER_COL  = 1
LEFT_MONTH_START = 1
LEFT_MONTH_END   = 17

# 右面板（CHN/NVN）的 section header 在 col 18
RIGHT_HEADER_COL  = 18
RIGHT_MONTH_START = 18
RIGHT_MONTH_END   = 35


def find_buffer_sheet_name(wb):
    for s in wb.sheetnames:
        if 'buffer' in s.lower():
            return s
    return None


def normalize_header(raw: str):
    """
    PJT section header → (div, region)，div 對齊 buffer sheet col A。
    '#22 IND'        → ('#22', 'IND')
    '#22CAB'         → ('#22', 'CAB')
    '#22 CHN'        → ('#22', 'CHN')
    '#22 NVN'        → ('#22', 'NVN')
    '#KIDS &#17 IND' → ('KIDS', 'IND')
    '# KIDS CHN'     → ('KIDS', 'CHN')
    '#KIDS  NVN'     → ('KIDS', 'NVN')
    '#22 SOLO'       → ('#22 solo', 'IND')
    '#32 IND SOLO'   → ('#32 solo', 'IND')
    '#32 IND'        → ('#32', 'IND')
    '# 33 NVN'       → ('#33', 'NVN')
    """
    s = raw.strip()
    s_up = s.upper()

    # 找産區（SOLO 優先，所以不提早 return）
    region = None
    for reg in REGIONS:
        if reg in s_up:
            region = reg
            break

    # KIDS
    if 'KIDS' in s_up:
        return 'KIDS', region

    # SOLO（有無 region 皆可，無 region 預設 IND）
    if 'SOLO' in s_up:
        m_num = re.search(r'#\s*(\d+)', s)
        if m_num:
            return '#' + m_num.group(1) + ' solo', region if region else 'IND'
        return None, None

    # 其餘需有明確産區
    if region is None:
        return None, None

    # 一般 #NN（含 #32 non-solo → '#32'）
    m = re.search(r'#\s*(\d+)', s)
    if m:
        return '#' + m.group(1), region

    return None, None


def get_month_col_map(ws_week, col_start, col_end):
    """從 row 3 讀月份縮寫 → 欄號（限定欄位範圍）。"""
    col_map = {}
    for c in range(col_start, col_end + 1):
        h = ws_week.cell(3, c).value
        if not h or not isinstance(h, str):
            continue
        h_up = h.strip().upper()
        for abbr in MONTH_ABBRS:
            if abbr.upper() in h_up:
                col_map[abbr] = c
                break
    return col_map


def get_buffer_row_map(ws_week):
    """
    掃描兩個面板，建立 (div, region) → (BUFFER列號, 月份欄panel) mapping。
    panel = 'left'（IND/CAB, 月份 col 5-13）
            'right'（CHN/NVN, 月份 col 22-30）
    同 section 只取第一個 BUFFER 列。
    """
    buf_map = {}

    for header_col, panel in [(LEFT_HEADER_COL, 'left'), (RIGHT_HEADER_COL, 'right')]:
        current_div = None
        current_region = None
        for r in range(1, ws_week.max_row + 1):
            v = ws_week.cell(r, header_col).value
            if v is None:
                continue
            vs = str(v).strip()

            if '#' in vs:
                d, reg = normalize_header(vs)
                if d and reg:
                    current_div = d
                    current_region = reg

            if vs.upper().startswith('BUFFER') and current_div and current_region:
                key = (current_div, current_region)
                if key not in buf_map:
                    buf_map[key] = (r, panel)

    return buf_map


def normalize_co(co_raw: str) -> str:
    s = str(co_raw).strip().upper()
    if s in ('NVN', 'N.VIN', 'NVIN', '北越'):
        return 'NVN'
    return s if s in ('IND', 'CAB', 'CHN', 'NVN') else s


def fill_buffer_formulas(xlsx_path: str, week_sheet: str, dry_run: bool = False):
    xlsx_path = Path(xlsx_path)

    if not dry_run:
        backup = xlsx_path.parent / f"{xlsx_path.stem}_backup_{week_sheet}{xlsx_path.suffix}"
        shutil.copy2(xlsx_path, backup)
        print(f'[備份] {backup}')

    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)

    if week_sheet not in wb.sheetnames:
        sys.exit(f'[ERROR] 找不到 sheet {repr(week_sheet)}，目前 sheets：{wb.sheetnames}')

    buf_sheet_name = find_buffer_sheet_name(wb)
    if not buf_sheet_name:
        sys.exit('[ERROR] 找不到 buffer 工作表')

    ws_week = wb[week_sheet]
    ws_buf  = wb[buf_sheet_name]

    print(f'週次 sheet：{week_sheet}')
    print(f'Buffer sheet：{buf_sheet_name}')

    # 兩套月份欄位 mapping
    month_col_left  = get_month_col_map(ws_week, LEFT_MONTH_START,  LEFT_MONTH_END)   # IND/CAB
    month_col_right = get_month_col_map(ws_week, RIGHT_MONTH_START, RIGHT_MONTH_END)  # CHN/NVN

    def fmt(m): return {k: f"{v}({get_column_letter(v)})" for k,v in sorted(m.items())}
    print(f'左面板月份（IND/CAB）：{fmt(month_col_left)}')
    print(f'右面板月份（CHN/NVN）：{fmt(month_col_right)}')

    # (div, region) → (BUFFER列號, panel)
    buf_row_map = get_buffer_row_map(ws_week)
    print(f'\n偵測到 {len(buf_row_map)} 個 section：')
    for (div, reg), (row, panel) in sorted(buf_row_map.items()):
        mcol = month_col_left if panel == 'left' else month_col_right
        vals = {get_column_letter(c): ws_week.cell(row, c).value
                for c in mcol.values()
                if ws_week.cell(row, c).value is not None}
        print(f'  {div:15s} {reg:3s}  [{panel:5s}] row {row:3d}  {vals}')

    # 填公式
    print()
    filled = zeroed = skipped = 0

    for r in range(1, 300):
        div_v  = ws_buf.cell(r, 1).value
        co_v   = ws_buf.cell(r, 2).value
        mon_v  = ws_buf.cell(r, 3).value
        item_v = ws_buf.cell(r, 5).value

        if not all([div_v, co_v, mon_v, item_v]):
            continue
        if str(item_v).strip() != 'Buffer':
            continue

        div_str = str(div_v).strip()
        co_str  = normalize_co(co_v)
        mon_raw = str(mon_v).strip()
        m_abbr  = re.search(r'([A-Za-z]{3})\s*$', mon_raw)
        mon_str = m_abbr.group(1).capitalize() if m_abbr else mon_raw.capitalize()

        key = (div_str, co_str)
        entry = buf_row_map.get(key)

        if entry is None:
            print(f'  零填 buf row {r:3d}  {div_str:15s} {co_str} {mon_str:3s}  →  0')
            if not dry_run:
                ws_buf.cell(r, 4).value = 0
            zeroed += 1
            continue

        buf_row, panel = entry
        month_col = month_col_left if panel == 'left' else month_col_right
        col_num = month_col.get(mon_str)

        if col_num is None:
            print(f'  [無月份] {div_str:15s} {co_str} {mon_str} — 週次 sheet 無此月欄')
            skipped += 1
            continue

        col_letter = get_column_letter(col_num)
        formula = f"='{week_sheet}'!{col_letter}{buf_row}"
        print(f'  寫入 buf row {r:3d}  {div_str:15s} {co_str} {mon_str:3s}  →  {formula}')

        if not dry_run:
            ws_buf.cell(r, 4).value = formula
        filled += 1

    if not dry_run:
        wb.save(str(xlsx_path))
        print(f'\n[完成] 填入 {filled} 個公式，補 0 {zeroed} 個，跳過 {skipped} 個')
    else:
        print(f'\n[Dry-run] 預計填入 {filled} 個公式，補 0 {zeroed} 個，跳過 {skipped} 個')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('xlsx', help='PJT xlsx 路徑')
    parser.add_argument('--week', required=True, help='週次 sheet 名稱，例如 07-14')
    parser.add_argument('--dry-run', action='store_true', help='只列出預計操作，不寫入')
    args = parser.parse_args()
    fill_buffer_formulas(args.xlsx, args.week, args.dry_run)
