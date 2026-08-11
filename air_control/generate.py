#!/usr/bin/env python3
"""由進單資料自動產生「大貨主副料空/陸運控管表」（含反黃區）。

用法:
    python3 generate.py orders_sample.csv
    python3 generate.py orders_sample.csv --outdir out --today 2026-08-11

輸出:
    控管表.xlsx   反黃已上色、凍結標題列、可直接發信的版本
    控管表.csv    給 BI / 其他系統再吃的純資料
    催辦.md       今天需要發問或決策的清單（貼進 mail 就能寄）
"""

from __future__ import annotations

import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rules import QTY_THRESHOLD, QUOTED_RATE, Row, apply_rules, needs_escalation

COLUMNS = [
    "產區", "DIV", "STYLE", "料別", "成衣件數", "面料數量_yds",
    "運輸方式", "運送數量_yds", "CRFP", "成衣EXP",
    "空陸運ETC", "空陸運IN_FTY", "海運IN_FTY", "報價已含",
    "空運單價", "實打", "超標金額_USD",
    "決策時間點", "狀態", "決策人員", "主要確認", "結論", "警示",
]

YELLOW = PatternFill("solid", fgColor="FFFF00")
PINK = PatternFill("solid", fgColor="FFD6E7")
GREY = PatternFill("solid", fgColor="D9D9D9")
HEADER = PatternFill("solid", fgColor="1F4E79")
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def parse_date(value: str) -> date | None:
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m/%d"):
        try:
            parsed = datetime.strptime(value, fmt)
        except ValueError:
            continue
        if fmt == "%m/%d":
            parsed = parsed.replace(year=date.today().year)
        return parsed.date()
    raise ValueError(f"看不懂的日期格式: {value!r}")


def parse_number(value: str) -> float | None:
    value = (value or "").strip().replace(",", "").replace("$", "")
    return float(value) if value else None


def load_rows(path: Path) -> list[Row]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        raw_rows = list(csv.DictReader(handle))

    rows: list[Row] = []
    for lineno, raw in enumerate(raw_rows, start=2):
        try:
            qty = parse_number(raw.get("成衣件數", ""))
            rows.append(Row(
                產區=raw.get("產區", "").strip(),
                DIV=raw.get("DIV", "").strip(),
                STYLE=raw.get("STYLE", "").strip(),
                料別=raw.get("料別", "").strip(),
                成衣件數=int(qty) if qty is not None else None,
                面料數量_yds=parse_number(raw.get("面料數量_yds", "")),
                運輸方式=raw.get("運輸方式", "").strip(),
                運送數量_yds=parse_number(raw.get("運送數量_yds", "")),
                CRFP=parse_date(raw.get("CRFP", "")),
                成衣EXP=parse_date(raw.get("成衣EXP", "")),
                空陸運ETC=parse_date(raw.get("空陸運ETC", "")),
                空陸運IN_FTY=parse_date(raw.get("空陸運IN_FTY", "")),
                海運IN_FTY=parse_date(raw.get("海運IN_FTY", "")),
                報價已含=raw.get("報價已含", "").strip(),
                空運單價=parse_number(raw.get("空運單價", "")),
                實打=parse_number(raw.get("實打", "")),
            ))
        except ValueError as exc:
            sys.exit(f"第 {lineno} 行資料有問題：{exc}")
    return rows


def enrich(rows: list[Row], today: date) -> list[Row]:
    """配布沿用同 STYLE 主布的成衣件數，再套規則。"""
    style_qty: dict[str, int] = {}
    for row in rows:
        if row.成衣件數 and row.STYLE not in style_qty:
            style_qty[row.STYLE] = row.成衣件數
    for row in rows:
        apply_rules(row, today, group_qty=style_qty.get(row.STYLE))
    return rows


def cell_value(row: Row, column: str):
    value = getattr(row, column)
    if column == "警示":
        return " / ".join(value)
    return value


def write_csv(rows: list[Row], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow([
                v.isoformat() if isinstance(v := cell_value(row, c), date) else v
                for c in COLUMNS
            ])


def write_xlsx(rows: list[Row], path: Path, today: date) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "控管表"

    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.fill = HEADER
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row in rows:
        sheet.append([cell_value(row, c) for c in COLUMNS])
        excel_row = sheet.max_row
        urgent = needs_escalation(row, today)
        for idx in range(1, len(COLUMNS) + 1):
            cell = sheet.cell(row=excel_row, column=idx)
            cell.border = BORDER
            if isinstance(cell.value, date):
                cell.number_format = "yyyy/mm/dd"
            if row.反黃:
                cell.fill = PINK if urgent else YELLOW
            elif row.運輸方式 == "海運":
                cell.fill = GREY
        if urgent:
            sheet.cell(row=excel_row, column=COLUMNS.index("狀態") + 1).font = Font(
                bold=True, color="C00000"
            )

    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"
    widths = {"STYLE": 18, "主要確認": 22, "結論": 26, "警示": 34, "決策人員": 14}
    for idx, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = widths.get(column, 13)

    workbook.save(path)


def write_escalation(rows: list[Row], path: Path, today: date) -> list[Row]:
    urgent = sorted(
        (r for r in rows if needs_escalation(r, today)),
        key=lambda r: r.決策時間點,
    )
    lines = [
        f"# 空/陸運決策催辦清單（{today:%Y-%m-%d}）",
        "",
        f"判斷條件：成衣量 > {QTY_THRESHOLD:,} 件，且決策日（ETC 前兩週）已到或剩 3 天內。",
        "",
    ]
    if not urgent:
        lines.append("今日無待決策項目。")
    else:
        lines.append("| STYLE | 料別 | 產區 | 成衣件數 | 運輸方式 | 決策日 | 狀態 | 決策人員 | 需確認 |")
        lines.append("|---|---|---|---|---|---|---|---|---|")
        for row in urgent:
            lines.append(
                f"| {row.STYLE} | {row.料別} | {row.產區} | {row.成衣件數 or '-':,} "
                f"| {row.運輸方式} | {row.決策時間點:%m/%d} | {row.狀態} "
                f"| {row.決策人員} | {row.主要確認} |"
                if isinstance(row.成衣件數, int)
                else f"| {row.STYLE} | {row.料別} | {row.產區} | - "
                f"| {row.運輸方式} | {row.決策時間點:%m/%d} | {row.狀態} "
                f"| {row.決策人員} | {row.主要確認} |"
            )

    over = [r for r in rows if r.超標金額_USD]
    if over:
        total = sum(r.超標金額_USD for r in over)
        lines += [
            "",
            f"## 超出報價標準 ${QUOTED_RATE}/實打 的空運（合計 US${total:,.0f}）",
            "",
            "| STYLE | 料別 | 單價 | 實打 | 超標金額 |",
            "|---|---|---|---|---|",
        ]
        for row in sorted(over, key=lambda r: -r.超標金額_USD):
            lines.append(
                f"| {row.STYLE} | {row.料別} | ${row.空運單價:.2f} "
                f"| {row.實打:,.0f} | ${row.超標金額_USD:,.0f} |"
            )

    gaps = [r for r in rows if r.警示]
    if gaps:
        lines += ["", "## 資料不全，請業務/採購補齊", ""]
        for row in gaps:
            lines.append(f"- {row.STYLE} {row.料別}：{' / '.join(row.警示)}")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return urgent


def main() -> None:
    parser = argparse.ArgumentParser(description="產生大貨空/陸運控管表")
    parser.add_argument("source", type=Path, help="進單資料 CSV")
    parser.add_argument("--outdir", type=Path, default=Path("out"))
    parser.add_argument("--today", type=parse_date, default=None,
                        help="模擬基準日 (YYYY-MM-DD)，預設今天")
    args = parser.parse_args()

    today = args.today or date.today()
    args.outdir.mkdir(parents=True, exist_ok=True)

    rows = enrich(load_rows(args.source), today)
    write_csv(rows, args.outdir / "控管表.csv")
    write_xlsx(rows, args.outdir / "控管表.xlsx", today)
    urgent = write_escalation(rows, args.outdir / "催辦.md", today)

    highlighted = sum(1 for r in rows if r.反黃)
    print(f"基準日 {today:%Y-%m-%d}｜共 {len(rows)} 筆，反黃 {highlighted} 筆，今日催辦 {len(urgent)} 筆")
    print(f"輸出於 {args.outdir.resolve()}")


if __name__ == "__main__":
    main()
